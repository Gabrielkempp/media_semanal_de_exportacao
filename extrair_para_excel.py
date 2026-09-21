"""
Rotina semanal: baixa a divulgação atual do MDIC, arquiva todos os arquivos do site e atualiza o Excel
com o histórico de todas as divulgações numa tabela única (aba DADOS), mais os textos (aba TEXTOS).

    python extrair_para_excel.py                -> baixa a divulgação atual e atualiza o Excel
    python extrair_para_excel.py --reprocessar  -> reconstrói o Excel a partir de saida/brutos, sem baixar nada

A pasta saida/brutos é a fonte da verdade: o Excel pode ser sempre reconstruído a partir dela.
Rodar mais de uma vez na mesma semana não duplica nada.
"""
import argparse
import logging
import os
import sys
from datetime import datetime
from pathlib import Path

import pandas as pd
import requests
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from src import principais_resultados as pr

PASTA_SAIDA = Path(__file__).parent / 'saida'
PASTA_BRUTOS = PASTA_SAIDA / 'brutos'
ARQUIVO_EXCEL = PASTA_SAIDA / 'balanca_semanal_mdic.xlsx'
ARQUIVO_LOG = PASTA_SAIDA / 'extracao.log'

def como_usar():
    """Regras da aba LEIA-ME, montadas a partir do que está configurado para captura."""
    fluxos = ' e '.join(pr.FLUXOS_CAPTURADOS).lower()
    if pr.PRODUTOS_CAPTURADOS is None:
        escopo = (f"Cada linha de DADOS é um produto de {fluxos} numa divulgação semanal do MDIC. "
                  f"Estão aqui todos os produtos publicados, então setor e total são a soma dos produtos.")
        recorte = ("Some VALOR_* e PESO_* à vontade dentro de uma mesma divulgação: os produtos somam o setor "
                   "e o total.")
    else:
        escopo = (f"Cada linha de DADOS é um produto de {fluxos} numa divulgação semanal do MDIC. "
                  f"Só são gravados estes produtos: {'; '.join(pr.PRODUTOS_CAPTURADOS)}.")
        recorte = ("ATENÇÃO: esta tabela é um recorte, não a balança inteira. A soma das linhas NÃO é o total "
                   "do setor nem do Brasil, e não dá para calcular saldo nem corrente de comércio. A divulgação "
                   "completa fica guardada em saida/brutos.")
    return [
        escopo,
        recorte,
        "Cada divulgação acumula o mês até DATA_REFERENCIA. Situação atual: filtre MAIS_RECENTE = 1. "
        "Comparar meses: filtre ULTIMA_DO_MES = 1. (1 = sim, 0 = não)",
        "Não some VALOR_ACUMULADO_MES_USD de divulgações diferentes do mesmo mês; para somar semanas use "
        "VALOR_SEMANA_USD.",
        "Para comparar períodos use as colunas MEDIA_DIARIA_*: cada período tem um número diferente de dias úteis.",
        "Variação de um conjunto de produtos = soma de MEDIA_DIARIA_USD ÷ soma de MEDIA_DIARIA_ANO_ANTERIOR_USD "
        "− 1. Nunca some nem tire média de VARIACAO_*.",
        "Valores em US$ e toneladas. Variações em fração (0,285 = 28,5%). O ano anterior é o mesmo mês completo, "
        "como publicado pelo MDIC na mesma planilha.",
        "Fonte: https://balanca.economia.gov.br/balanca/pg_principal_bc/principais_resultados.html "
        "(arquivos originais de cada semana em saida/brutos).",
    ]

log = logging.getLogger('mdic')


# ---------------------------------------------------------------- download e arquivamento

def salvar_brutos(pasta, arquivos):
    pasta.mkdir(parents=True, exist_ok=True)
    for nome, conteudo in arquivos.items():
        (pasta / nome).write_bytes(conteudo)
    return pasta


def ler_brutos(pasta):
    return {arquivo.name: arquivo.read_bytes() for arquivo in pasta.iterdir() if arquivo.is_file()}


def baixar_divulgacao_atual():
    """Baixa a página e todos os arquivos linkados nela e arquiva em saida/brutos/<ano>-<mes>_semana<n>."""
    html = pr.baixar(pr.URL_PAGINA)
    try:
        parser = pr.analisar_pagina(html.decode('utf-8'))
        pub = pr.publicacao_da_pagina(parser)
    except Exception:
        # Guarda a página para análise: o layout do site provavelmente mudou
        salvar_brutos(PASTA_BRUTOS / f'_falha_{datetime.now():%Y-%m-%d_%H%M%S}', {pr.ARQ_PAGINA: html})
        raise
    log.info(f"Divulgação no site: {pub.descricao} (publicada em {pub.data_publicacao:%d/%m/%Y})")

    arquivos = {pr.ARQ_PAGINA: html}
    for nome, url in pr.links_da_pagina(parser).items():
        try:
            arquivos[nome] = pr.baixar(url)
        except requests.RequestException as erro:
            log.warning(f"Não foi possível baixar {url}: {erro}")
    pasta = salvar_brutos(PASTA_BRUTOS / pub.pasta, arquivos)
    log.info(f"{len(arquivos)} arquivos arquivados em {pasta}")
    return pub


# ---------------------------------------------------------------- Excel

def ids_das_linhas(df):
    return pr.id_publicacao(df['ANO'], df['MES'], df['SEMANA'])


def ler_excel_existente():
    """Tabelas atuais do Excel, ou None se não existir ou se a estrutura mudou (aí tudo é reprocessado)."""
    if not ARQUIVO_EXCEL.exists():
        return None
    existentes = pd.read_excel(ARQUIVO_EXCEL, sheet_name=None)
    for tabela, colunas in pr.ESQUEMA.items():
        if tabela not in existentes or list(existentes[tabela].columns) != colunas:
            log.info("Estrutura do Excel diferente da atual; reconstruindo a partir dos arquivos brutos")
            return None
    return {tabela: existentes[tabela] for tabela in pr.ESQUEMA}


def combinar(existentes, novas):
    """Junta o histórico com as divulgações processadas agora (que substituem as de mesma semana)."""
    ids_novos = {int(ids_das_linhas(n['DADOS']).iloc[0]) for n in novas}
    tabelas = {}
    for tabela in pr.ESQUEMA:
        partes = [n[tabela] for n in novas]
        if tabela in existentes:
            antigo = existentes[tabela]
            partes.insert(0, antigo[~ids_das_linhas(antigo).isin(ids_novos)])
        tabelas[tabela] = pd.concat(partes, ignore_index=True).sort_values('DATA_REFERENCIA', kind='stable')
    dados, textos = pr.calcular_colunas_historicas(tabelas['DADOS'], tabelas['TEXTOS'])
    return {'DADOS': dados.reset_index(drop=True), 'TEXTOS': textos.reset_index(drop=True)}


def formato_numero(coluna):
    if coluna.startswith('VARIACAO_'):
        return '0.0%'
    if coluna in pr.COLUNAS_DATA:
        return 'DD/MM/YYYY'
    if coluna.startswith('PRECO_'):
        return '#,##0.00'
    if coluna.endswith('_USD'):
        return '#,##0'
    if coluna.endswith('_TON'):
        return '#,##0.0'
    return None


def formatar_aba(ws, df, nome):
    """A aba vira uma Tabela do Excel com o mesmo nome, que é o que o Power BI lista ao importar."""
    if len(df):
        tabela = Table(displayName=nome, ref=f"A1:{get_column_letter(len(df.columns))}{len(df) + 1}")
        tabela.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
        ws.add_table(tabela)
    ws.freeze_panes = 'A2'
    for i, coluna in enumerate(df.columns, start=1):
        letra = get_column_letter(i)
        if coluna == 'TEXTO':
            ws.column_dimensions[letra].width = 110
            for (celula,) in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                celula.alignment = Alignment(wrap_text=True, vertical='top')
            continue
        formato = formato_numero(coluna)
        if formato:
            for (celula,) in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                celula.number_format = formato
        maior_valor = df[coluna].head(2000).astype(str).str.len().max() if len(df) else 0
        ws.column_dimensions[letra].width = min(max(len(coluna), maior_valor) + 3, 60)


def escrever_leia_me(ws):
    ws.append(['Balança comercial semanal (MDIC): como usar'])
    ws['A1'].font = Font(bold=True, size=13)
    for i, regra in enumerate(como_usar(), start=1):
        ws.append([f"{i}. {regra}"])
    ws.append([])
    ws.append(['TABELA', 'COLUNA', 'DESCRIÇÃO'])
    for celula in ws[ws.max_row]:
        celula.font = Font(bold=True)
    for tabela, colunas in pr.DESCRICOES.items():
        for coluna, descricao in colunas.items():
            ws.append([tabela, coluna, descricao])
    ws.column_dimensions['A'].width = 10
    ws.column_dimensions['B'].width = 34
    ws.column_dimensions['C'].width = 130


def gravar_excel(tabelas):
    # Grava num arquivo temporário e troca no fim: uma falha no meio nunca corrompe o Excel atual
    temporario = ARQUIVO_EXCEL.with_name(ARQUIVO_EXCEL.stem + '.tmp.xlsx')
    with pd.ExcelWriter(temporario, engine='openpyxl') as writer:
        for nome, df in tabelas.items():
            df.to_excel(writer, sheet_name=nome, index=False)
            formatar_aba(writer.sheets[nome], df, nome)
        escrever_leia_me(writer.book.create_sheet('LEIA-ME'))
    try:
        os.replace(temporario, ARQUIVO_EXCEL)
    except PermissionError:
        temporario.unlink(missing_ok=True)
        raise RuntimeError(f"{ARQUIVO_EXCEL.name} está aberto em outro programa; feche e rode novamente "
                           f"(os arquivos brutos já foram salvos, nada se perde)") from None


def sincronizar(reprocessar=False, obrigatorias=()):
    """
    Processa as divulgações arquivadas que ainda não estão no Excel (e as `obrigatorias`) e grava.
    Retorna os IDs que falharam.
    """
    existentes = None if reprocessar else ler_excel_existente()
    ids_existentes = set(ids_das_linhas(existentes['DADOS'])) if existentes else set()

    novas, falhas = [], []
    for pasta in sorted(PASTA_BRUTOS.glob('*_semana*')):
        id_pub = pr.id_da_pasta(pasta.name)
        if id_pub is None or (id_pub in ids_existentes and id_pub not in obrigatorias):
            continue
        try:
            tabelas = pr.processar_divulgacao(ler_brutos(pasta))
            id_conteudo = int(ids_das_linhas(tabelas['DADOS']).iloc[0])
            if id_conteudo != id_pub:
                raise ValueError(f"A pasta {pasta.name} contém a divulgação {id_conteudo}")
            novas.append(tabelas)
            log.info(f"Processada: {pasta.name}")
        except Exception:
            log.exception(f"Falha ao processar {pasta.name}")
            falhas.append(id_pub)

    if not novas:
        log.info("Nenhuma divulgação nova para gravar")
        return falhas
    tabelas = combinar(existentes or {}, novas)
    gravar_excel(tabelas)
    log.info("Excel atualizado: " + ", ".join(f"{nome}={len(df)} linhas" for nome, df in tabelas.items()))
    return falhas


# ---------------------------------------------------------------- execução

def configurar_log():
    PASTA_SAIDA.mkdir(exist_ok=True)
    logging.basicConfig(
        level=logging.INFO,
        format='%(asctime)s %(levelname)s %(message)s',
        handlers=[logging.FileHandler(ARQUIVO_LOG, encoding='utf-8'), logging.StreamHandler()],
    )


def main():
    argumentos = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    argumentos.add_argument('--reprocessar', action='store_true',
                            help='reconstrói o Excel a partir de todos os arquivos brutos, sem baixar nada')
    argumentos.add_argument('--ddl', action='store_true',
                            help='imprime o CREATE TABLE das tabelas (para criar no banco) e sai')
    args = argumentos.parse_args()

    if args.ddl:
        for tabela in pr.ESQUEMA:
            print(pr.ddl_tabela(tabela), end='\n\n')
        return 0

    configurar_log()

    try:
        obrigatorias = set()
        if not args.reprocessar:
            obrigatorias.add(baixar_divulgacao_atual().id)
        falhas = sincronizar(args.reprocessar, obrigatorias)
    except Exception:
        log.exception("Execução interrompida")
        return 1

    # Falha na divulgação atual (ou em qualquer uma, ao reprocessar) deixa a tarefa agendada marcada como erro
    if (obrigatorias & set(falhas)) or (args.reprocessar and falhas):
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
